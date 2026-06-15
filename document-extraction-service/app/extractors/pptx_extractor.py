import base64
import io

import structlog
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from app.extractors.base import BaseExtractor
from app.models.semantic import ImageData, SemanticExtractionResult
from app.utils.text import clean_cell_value, normalize_whitespace

logger = structlog.get_logger()


class PptxExtractor(BaseExtractor):
    supported_extensions = ["pptx"]

    def extract(self, file_bytes: bytes, filename: str) -> SemanticExtractionResult:
        prs = Presentation(io.BytesIO(file_bytes))

        blocks = []
        heading_stack: dict[int, str] = {}
        block_index = 0

        title = ""

        for slide_number, slide in enumerate(prs.slides, start=1):
            # Slide title as heading
            slide_title = ""
            if slide.shapes.title and slide.shapes.title.has_text_frame:
                slide_title = normalize_whitespace(slide.shapes.title.text_frame.text)

            if slide_title:
                if not title:
                    title = slide_title
                blocks.append(
                    self._make_heading_block(
                        block_index, slide_title, 1, heading_stack, slide_number=slide_number
                    )
                )
                block_index += 1

            # Process shapes
            for shape in slide.shapes:
                # Skip the title shape (already handled)
                if shape == slide.shapes.title:
                    continue

                # Pictures
                if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                    try:
                        img_blob = shape.image.blob
                        img_format = shape.image.content_type.split("/")[-1] if "/" in shape.image.content_type else "png"
                        # Convert EMUs to pixels (1 inch = 914400 EMUs, 96 DPI)
                        px_width = round(shape.width / 9525) if shape.width else None
                        px_height = round(shape.height / 9525) if shape.height else None
                        alt_text = shape.name or ""
                        b64 = base64.b64encode(img_blob).decode("ascii")
                        image_data = ImageData(
                            base64=b64, format=img_format,
                            width=px_width, height=px_height, alt_text=alt_text,
                        )
                        blocks.append(self._make_image_block(
                            block_index, image_data, heading_stack, slide_number=slide_number
                        ))
                        block_index += 1
                    except Exception as e:
                        logger.warning("pptx_image_failed", shape=shape.name, error=str(e))

                # Tables
                elif shape.has_table:
                    headers, rows = self._parse_table(shape.table)
                    if headers or rows:
                        blocks.append(
                            self._make_table_block(
                                block_index, headers, rows, heading_stack, slide_number=slide_number
                            )
                        )
                        block_index += 1

                # Text frames
                elif shape.has_text_frame:
                    for paragraph in shape.text_frame.paragraphs:
                        text = normalize_whitespace(paragraph.text)
                        if not text:
                            continue

                        # Check if paragraph level suggests a sub-heading
                        if paragraph.level == 0 and self._is_bold(paragraph):
                            blocks.append(
                                self._make_heading_block(
                                    block_index, text, 2, heading_stack, slide_number=slide_number
                                )
                            )
                        else:
                            blocks.append(
                                self._make_paragraph_block(
                                    block_index, text, heading_stack, slide_number=slide_number
                                )
                            )
                        block_index += 1

        metadata = self._create_metadata(
            filename, title=title, page_count=len(prs.slides)
        )
        raw_content = self._build_raw_content(blocks)

        return SemanticExtractionResult(metadata=metadata, blocks=blocks, raw_content=raw_content)

    def _parse_table(self, table) -> tuple[list[str], list[list[str]]]:
        all_rows = []
        for row in table.rows:
            cells = [clean_cell_value(cell.text) for cell in row.cells]
            all_rows.append(cells)

        if not all_rows:
            return [], []

        return all_rows[0], all_rows[1:]

    def _is_bold(self, paragraph) -> bool:
        """Check if the first run of a paragraph is bold."""
        if paragraph.runs:
            return paragraph.runs[0].font.bold is True
        return False
